"""Builds the downloadable Excel workbook (UC-5).

Excel is the output format because it is what Saudi university administrative
workflows already use, and results must drop into that workflow without a
conversion step (SS 2.24).

The feedback column is written right-to-left with an explicit reading order so
Arabic renders correctly in Excel rather than appearing reversed or
left-aligned (SS 2.23).
"""
import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)

RTL = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)
LTR = Alignment(horizontal="left", vertical="top", wrap_text=True, readingOrder=1)
CENTER = Alignment(horizontal="center", vertical="top")

COLUMNS = [
    ("Student ID", 18),
    ("Score", 10),
    ("Feedback (Arabic)", 80),
    ("Note", 28),
]


def format_score(value):
    """5.0 -> "5", 2.5 -> "2.5"."""
    return "%g" % value


def format_breakdown(breakdown):
    """"Definition 3/4 · Example 4/6", for the export and the results page."""
    return " · ".join(
        p["name"] + " " + format_score(p["score"]) + "/" + format_score(p["marks"])
        for p in breakdown
    )


def _header(sheet, columns):
    for index, (title, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def _score(cell, value):
    cell.value = value if value is not None else "-"
    cell.alignment = CENTER
    if value is not None:
        cell.number_format = "0.##"


def build_workbook(assignment, results, questions=()):
    """Return the workbook as bytes.

    `assignment` is the assignment row; `results` is a sequence of rows with
    student_id, score, feedback_ar and error. For an exam, `questions` holds
    its question rows and each result also carries question_id.
    """
    if questions:
        return _build_exam_workbook(assignment, results, questions)

    # A "Parts" column only when the question was marked in parts.
    columns = list(COLUMNS)
    in_parts = any(row.get("breakdown") for row in results)
    if in_parts:
        columns.append(("Parts", 40))

    book = Workbook()
    sheet = book.active
    sheet.title = "Results"

    _header(sheet, columns)

    for offset, row in enumerate(results, start=2):
        sheet.cell(row=offset, column=1, value=row["student_id"]).alignment = LTR

        score_cell = sheet.cell(row=offset, column=2)
        score_cell.value = row["score"] if row["score"] is not None else "-"
        score_cell.alignment = CENTER
        if row["score"] is not None:
            score_cell.number_format = "0.##"

        feedback = sheet.cell(row=offset, column=3, value=row["feedback_ar"] or "")
        feedback.alignment = RTL

        note = sheet.cell(row=offset, column=4, value=row["error"] or "")
        note.alignment = LTR

        if in_parts:
            sheet.cell(row=offset, column=5,
                       value=format_breakdown(row.get("breakdown") or [])).alignment = LTR

    _add_metadata_sheet(book, assignment, results)

    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def _build_exam_workbook(assignment, results, questions):
    """One row per student with a score per question and a total, then a
    second sheet with every answer's feedback."""
    from .results import group_by_student

    students = group_by_student(results)
    book = Workbook()

    sheet = book.active
    sheet.title = "Results"
    columns = [("Student ID", 18)]
    columns += [("Q" + str(q["position"]) + " (/" + "%g" % q["max_score"] + ")", 12)
                for q in questions]
    columns.append(("Total (/" + "%g" % assignment["max_score"] + ")", 14))
    _header(sheet, columns)
    for offset, student in enumerate(students, start=2):
        sheet.cell(row=offset, column=1, value=student["student_id"]).alignment = LTR
        for index, answer in enumerate(student["answers"], start=2):
            _score(sheet.cell(row=offset, column=index), answer["score"])
        total = sheet.cell(row=offset, column=len(columns), value=None)
        _score(total, student["total"])
        total.font = Font(bold=True)

    detail = book.create_sheet("Feedback")
    _header(detail, [("Student ID", 18), ("Question", 28), ("Score", 10),
                     ("Feedback (Arabic)", 80), ("Note", 28), ("Parts", 40)])
    offset = 2
    for student in students:
        for question, answer in zip(questions, student["answers"]):
            detail.cell(row=offset, column=1, value=student["student_id"]).alignment = LTR
            detail.cell(row=offset, column=2,
                        value="Q" + str(question["position"]) + " " + question["title"]
                        ).alignment = LTR
            _score(detail.cell(row=offset, column=3), answer["score"])
            detail.cell(row=offset, column=4, value=answer["feedback_ar"] or "").alignment = RTL
            detail.cell(row=offset, column=5, value=answer["error"] or "").alignment = LTR
            detail.cell(row=offset, column=6,
                        value=format_breakdown(answer.get("breakdown") or [])).alignment = LTR
            offset += 1

    meta = _add_metadata_sheet(book, assignment, results)
    # Replace the single ideal answer and rubric with one entry per question.
    for row in range(meta.max_row, 0, -1):
        if meta.cell(row=row, column=1).value in ("Ideal answer", "Rubric"):
            meta.delete_rows(row)
    totals = [s["total"] for s in students if s["total"] is not None]
    meta.append(("Students", len(students)))
    meta.append(("Average total", round(sum(totals) / len(totals), 2) if totals else "-"))
    for q in questions:
        label = "Q" + str(q["position"])
        meta.append((label + " column", q["title"]))
        meta.append((label + " marks", q["max_score"]))
        meta.append((label + " ideal answer", q["ideal_answer"]))
        meta.append((label + " rubric", q["rubric"] or "-"))
        if q["parts"]:
            meta.append((label + " parts", "; ".join(
                p["name"] + " (" + format_score(p["marks"]) + ")" for p in json.loads(q["parts"])
            )))
    for row in meta.iter_rows(min_row=1):
        row[0].font = Font(bold=True)
        row[0].alignment = Alignment(vertical="top")
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def _add_metadata_sheet(book, assignment, results):
    """Record what produced these grades.

    Useful for the instructor, and it is also the provenance the project report
    needs when quoting measured results.
    """
    sheet = book.create_sheet("Assignment")
    scored = [r["score"] for r in results if r["score"] is not None]
    models = sorted({r["model_name"] for r in results if r["model_name"]})

    rows = [
        ("Assignment", assignment["title"]),
        ("Graded on", assignment["created_at"]),
        ("Maximum score", assignment["max_score"]),
        ("Responses", len(results)),
        ("Graded successfully", len(scored)),
        ("Average score", round(sum(scored) / len(scored), 2) if scored else "-"),
        ("Model", ", ".join(models) or "-"),
        ("Ideal answer", assignment["ideal_answer"]),
        ("Rubric", assignment["rubric"] or "-"),
    ]
    parts = assignment.get("parts") if hasattr(assignment, "get") else assignment["parts"]
    if parts:
        rows.append(("Parts", "; ".join(
            p["name"] + " (" + format_score(p["marks"]) + ")" for p in json.loads(parts)
        )))

    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 90
    for index, (label, value) in enumerate(rows, start=1):
        key = sheet.cell(row=index, column=1, value=label)
        key.font = Font(bold=True)
        key.alignment = Alignment(vertical="top")
        sheet.cell(row=index, column=2, value=value).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
    return sheet


def filename_for(assignment):
    """A safe download filename derived from the assignment title.

    The uploaded file's own extension is dropped first, so "midterm.csv" yields
    "qayyem-midterm.xlsx" rather than "qayyem-midterm_csv.xlsx".
    """
    title = assignment["title"] or "results"
    for extension in (".csv", ".xlsx"):
        if title.lower().endswith(extension):
            title = title[: -len(extension)]
            break

    base = "".join(
        ch if (ch.isalnum() or ch in " -_") else "_" for ch in title
    ).strip("_ ") or "results"
    return "qayyem-" + base.replace(" ", "-")[:60] + ".xlsx"
