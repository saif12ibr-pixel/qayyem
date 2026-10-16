"""End-to-end tests: upload, grade, view, download.

These run the real pipeline against the deterministic fake grader, so they cover
UC-2 through UC-5 without a network call.
"""
import io

import pytest
from openpyxl import load_workbook

from qayyem.db import get_db
from qayyem.grading.prompt import EMPTY_RESPONSE_FEEDBACK
from tests.conftest import INSTRUCTOR_EMAIL

IDEAL = "Photosynthesis converts light energy into chemical energy."


@pytest.fixture
def app(app):  # noqa: F811 - deliberately wraps the conftest fixture
    """Grade inline so assertions never race the worker thread."""
    app.config["GRADING_SYNC"] = True
    return app


def upload(client, csv_text, ideal=IDEAL, rubric="Accuracy and completeness.", max_score="10"):
    return client.post(
        "/upload",
        data={
            "file": (io.BytesIO(csv_text.encode("utf-8")), "responses.csv"),
            "ideal_answer": ideal,
            "rubric": rubric,
            "max_score": max_score,
        },
        content_type="multipart/form-data",
        follow_redirects=False,
    )


def test_full_upload_grade_download_flow(client, app, auth):
    auth.login()

    response = upload(
        client,
        "student_id,response\n"
        "S001,Photosynthesis converts light energy into chemical energy.\n"
        "S002,Plants are green.\n",
    )
    assert response.status_code == 302
    assignment_id = int(response.headers["Location"].rstrip("/").split("/")[-1])

    with app.app_context():
        row = get_db().execute(
            "SELECT status, progress_done, progress_total FROM assignment WHERE id = ?",
            (assignment_id,),
        ).fetchone()
        assert row["status"] == "complete"
        assert row["progress_done"] == row["progress_total"] == 2

    page = client.get("/results/" + str(assignment_id))
    assert page.status_code == 200
    assert b"S001" in page.data

    status = client.get("/results/" + str(assignment_id) + "/status").get_json()
    assert status["status"] == "complete"
    assert status["done"] == 2

    export = client.get("/results/" + str(assignment_id) + "/export.xlsx")
    assert export.status_code == 200
    assert "attachment" in export.headers["Content-Disposition"]

    book = load_workbook(io.BytesIO(export.data))
    sheet = book["Results"]
    assert sheet["A1"].value == "Student ID"
    assert {sheet.cell(row=r, column=1).value for r in (2, 3)} == {"S001", "S002"}


def test_matching_answer_scores_higher_than_unrelated_one(client, app, auth):
    auth.login()
    response = upload(
        client,
        "student_id,response\nS001," + IDEAL + "\nS002,totally unrelated gibberish\n",
    )
    assignment_id = int(response.headers["Location"].rstrip("/").split("/")[-1])

    with app.app_context():
        rows = get_db().execute(
            """SELECT sr.student_id, gr.score FROM student_response sr
               JOIN grading_result gr ON gr.response_id = sr.id
               WHERE sr.assignment_id = ? ORDER BY sr.student_id""",
            (assignment_id,),
        ).fetchall()
    scores = {r["student_id"]: r["score"] for r in rows}
    assert scores["S001"] > scores["S002"]


def test_empty_response_scores_zero_with_default_feedback(client, app, auth):
    """UC-3 alternate flow 2."""
    auth.login()
    response = upload(client, "student_id,response\nS001,\nS002,some answer\n")
    assignment_id = int(response.headers["Location"].rstrip("/").split("/")[-1])

    with app.app_context():
        row = get_db().execute(
            """SELECT gr.score, gr.feedback_ar FROM student_response sr
               JOIN grading_result gr ON gr.response_id = sr.id
               WHERE sr.assignment_id = ? AND sr.student_id = 'S001'""",
            (assignment_id,),
        ).fetchone()
    assert row["score"] == 0
    assert row["feedback_ar"] == EMPTY_RESPONSE_FEEDBACK


def test_batching_covers_every_student(client, app, auth):
    """Batch size is 3 in tests, so 7 responses exercise uneven final batch."""
    auth.login()
    body = "".join("S{0:03d},answer number {0}\n".format(i) for i in range(1, 8))
    response = upload(client, "student_id,response\n" + body)
    assignment_id = int(response.headers["Location"].rstrip("/").split("/")[-1])

    with app.app_context():
        count = get_db().execute(
            """SELECT COUNT(*) c FROM grading_result gr
               JOIN student_response sr ON sr.id = gr.response_id
               WHERE sr.assignment_id = ?""",
            (assignment_id,),
        ).fetchone()["c"]
    assert count == 7


def test_upload_rejects_missing_ideal_answer(client, auth):
    auth.login()
    response = upload(client, "student_id,response\nS001,x\n", ideal="")
    assert response.status_code == 400
    assert b"ideal answer" in response.data


def test_upload_rejects_bad_max_score(client, auth):
    auth.login()
    response = upload(client, "student_id,response\nS001,x\n", max_score="0")
    assert response.status_code == 400
    assert b"Maximum score" in response.data


def test_upload_rejects_malformed_file_with_message(client, auth):
    auth.login()
    response = upload(client, "name,grade\nAli,5\n")
    assert response.status_code == 400
    assert b"column" in response.data


def test_export_blocked_before_grading_completes(client, app, auth):
    """UC-5 alternate flow."""
    auth.login()
    with app.app_context():
        db = get_db()
        # Select the signed-in instructor explicitly. "LIMIT 1" without ORDER BY
        # returns whichever row the planner reaches first, which is not row 1.
        instructor_id = db.execute(
            "SELECT id FROM instructor WHERE email = ?", (INSTRUCTOR_EMAIL,)
        ).fetchone()["id"]
        cursor = db.execute(
            "INSERT INTO assignment (instructor_id, title, ideal_answer, max_score, status) "
            "VALUES (?, 'x.csv', 'ideal', 10, 'grading')",
            (instructor_id,),
        )
        db.commit()
        assignment_id = cursor.lastrowid

    response = client.get("/results/" + str(assignment_id) + "/export.xlsx")
    assert response.status_code == 302


def test_export_marks_assignment_as_downloaded(client, app, auth):
    auth.login()
    response = upload(client, "student_id,response\nS001,answer\n")
    assignment_id = int(response.headers["Location"].rstrip("/").split("/")[-1])

    with app.app_context():
        assert get_db().execute(
            "SELECT exported_at FROM assignment WHERE id = ?", (assignment_id,)
        ).fetchone()["exported_at"] is None

    client.get("/results/" + str(assignment_id) + "/export.xlsx")

    with app.app_context():
        assert get_db().execute(
            "SELECT exported_at FROM assignment WHERE id = ?", (assignment_id,)
        ).fetchone()["exported_at"] is not None


def test_rate_limit_stops_the_run_instead_of_amplifying_requests(client, app, auth):
    """A 429 means "stop sending requests".

    The batch fallback retries items individually, which on a rate limit would
    send an order of magnitude MORE requests into a provider already refusing
    us, burning a quota that is often a daily allowance. So a rate limit must
    end the run, not fan out.
    """
    from qayyem.grading import GraderRateLimitError
    from qayyem.grading.base import Grader
    from qayyem.db import get_db
    # Patch where the name is looked up: settings.py does
    # `from .grading.factory import build_grader`, so it holds its own
    # reference and patching the factory module would not reach it.
    import qayyem.settings as settings_module

    calls = {"n": 0}

    class RateLimitedGrader(Grader):
        name = "rate-limited-test-grader"

        def grade_batch(self, items, ideal_answer, rubric, max_score):
            calls["n"] += 1
            raise GraderRateLimitError("429 quota exceeded")

    original = settings_module.build_grader
    settings_module.build_grader = lambda *a, **k: RateLimitedGrader()
    try:
        auth.login()
        body = "".join("S{0:03d},answer {0}\n".format(i) for i in range(1, 8))
        response = upload(client, "student_id,response\n" + body)
        assignment_id = int(response.headers["Location"].rstrip("/").split("/")[-1])
    finally:
        settings_module.build_grader = original

    # Batch size is 3 in tests: 7 responses would be 3 batches, and fanning out
    # would be many more calls. One call then stop is the required behaviour.
    assert calls["n"] == 1

    with app.app_context():
        row = get_db().execute(
            "SELECT status, error_message FROM assignment WHERE id = ?", (assignment_id,)
        ).fetchone()
        assert row["status"] == "failed"
        assert "rate limit" in row["error_message"].lower()
        assert "kept" in row["error_message"]

        # The uploaded responses must survive so the run can be repeated.
        kept = get_db().execute(
            "SELECT COUNT(*) c FROM student_response WHERE assignment_id = ?",
            (assignment_id,),
        ).fetchone()["c"]
        assert kept == 7
