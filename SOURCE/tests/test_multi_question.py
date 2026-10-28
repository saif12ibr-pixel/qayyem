"""Multi-question exam files: one answer column per question, each graded
against its own ideal answer and worth its own marks."""
import io
import sqlite3

import pytest
from openpyxl import load_workbook

from qayyem.db import ensure_schema_current, get_db
from qayyem.parsing import ParseError, parse_exam

Q1_IDEAL = "Photosynthesis converts light energy into chemical energy."
Q2_IDEAL = "Mitochondria release energy through respiration."

EXAM_CSV = (
    "student_id,answer 1,answer 2,answer 3\n"
    "S001," + Q1_IDEAL + "," + Q2_IDEAL + ",water\n"
    "S002,plants are green,,water\n"
)


@pytest.fixture
def app(app):  # noqa: F811 - deliberately wraps the conftest fixture
    app.config["GRADING_SYNC"] = True
    return app


def upload_exam(client, csv_text=EXAM_CSV, max_score="10"):
    return client.post(
        "/upload",
        data={
            "file": (io.BytesIO(csv_text.encode("utf-8")), "exam.csv"),
            "rubric": "Accuracy.",
            "max_score": max_score,
        },
        content_type="multipart/form-data",
    )


def question_ids(app, assignment_id):
    with app.app_context():
        return [r["id"] for r in get_db().execute(
            "SELECT id FROM question WHERE assignment_id = ? ORDER BY position",
            (assignment_id,),
        )]


def setup_form(ids, marks=("5", "10", "15")):
    form = {}
    ideals = (Q1_IDEAL, Q2_IDEAL, "water")
    for qid, ideal, mark in zip(ids, ideals, marks):
        form["ideal_answer_" + str(qid)] = ideal
        form["rubric_" + str(qid)] = "Accuracy."
        form["max_score_" + str(qid)] = mark
    return form


def start_exam(client, app):
    response = upload_exam(client)
    assert response.status_code == 302
    assert "/questions" in response.headers["Location"]
    assignment_id = int(response.headers["Location"].split("/")[-2])
    ids = question_ids(app, assignment_id)
    response = client.post(
        "/upload/" + str(assignment_id) + "/questions",
        data=setup_form(ids),
        content_type="multipart/form-data",
    )
    assert response.status_code == 302
    return assignment_id


# --- parsing ---------------------------------------------------------------

def test_parse_exam_returns_every_answer_column():
    exam = parse_exam("exam.csv", EXAM_CSV.encode("utf-8"))
    assert exam["questions"] == ["answer 1", "answer 2", "answer 3"]
    assert exam["students"][1]["answers"] == ["plants are green", "", "water"]


def test_parse_exam_treats_a_single_response_column_as_one_question():
    exam = parse_exam("a.csv", b"student_id,response\nS001,x\n")
    assert exam["questions"] == ["response"]


def test_parse_exam_limits_question_count():
    headers = ["student_id"] + ["answer {0}".format(i) for i in range(1, 30)]
    data = (",".join(headers) + "\nS001," + ",".join("x" * 28) + "\n").encode("utf-8")
    with pytest.raises(ParseError, match="limit of 20 questions"):
        parse_exam("exam.csv", data)


# --- flow ------------------------------------------------------------------

def test_exam_upload_asks_for_each_question(client, app, auth):
    auth.login()
    response = upload_exam(client)
    page = client.get(response.headers["Location"])
    assert page.status_code == 200
    for title in (b"answer 1", b"answer 2", b"answer 3"):
        assert title in page.data
    # The results page sends an unconfigured exam back to setup.
    assignment_id = response.headers["Location"].split("/")[-2]
    redirect = client.get("/results/" + assignment_id)
    assert "/questions" in redirect.headers["Location"]


def test_each_question_has_its_own_marks_and_ideal_answer(client, app, auth):
    auth.login()
    assignment_id = start_exam(client, app)

    with app.app_context():
        db = get_db()
        assignment = db.execute(
            "SELECT status, max_score FROM assignment WHERE id = ?", (assignment_id,)
        ).fetchone()
        rows = db.execute(
            """SELECT sr.student_id, q.position, q.max_score, gr.score
               FROM student_response sr
               JOIN question q ON q.id = sr.question_id
               JOIN grading_result gr ON gr.response_id = sr.id
               WHERE sr.assignment_id = ?""",
            (assignment_id,),
        ).fetchall()

    assert assignment["status"] == "complete"
    assert assignment["max_score"] == 30  # 5 + 10 + 15
    scores = {(r["student_id"], r["position"]): r["score"] for r in rows}
    # S001 matches each question's own ideal answer, so gets full marks on each.
    assert scores[("S001", 1)] == 5
    assert scores[("S001", 2)] == 10
    assert scores[("S001", 3)] == 15
    # S002 left question 2 blank: zero without an AI call.
    assert scores[("S002", 2)] == 0


def test_results_and_export_show_per_question_scores_and_total(client, app, auth):
    auth.login()
    assignment_id = start_exam(client, app)

    page = client.get("/results/" + str(assignment_id))
    assert page.status_code == 200
    assert b"Total" in page.data and b"Q3" in page.data

    export = client.get("/results/" + str(assignment_id) + "/export.xlsx")
    book = load_workbook(io.BytesIO(export.data))
    sheet = book["Results"]
    headers = [c.value for c in sheet[1]]
    assert headers == ["Student ID", "Q1 (/5)", "Q2 (/10)", "Q3 (/15)", "Total (/30)"]
    assert [c.value for c in sheet[2]] == ["S001", 5, 10, 15, 30]
    assert book["Feedback"].max_row == 1 + 2 * 3


def test_setup_rejects_a_question_without_an_ideal_answer(client, app, auth):
    auth.login()
    response = upload_exam(client)
    assignment_id = int(response.headers["Location"].split("/")[-2])
    ids = question_ids(app, assignment_id)
    form = setup_form(ids)
    form["ideal_answer_" + str(ids[1])] = ""

    response = client.post(
        "/upload/" + str(assignment_id) + "/questions",
        data=form,
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert b"Question 2" in response.data
    with app.app_context():
        status = get_db().execute(
            "SELECT status FROM assignment WHERE id = ?", (assignment_id,)
        ).fetchone()["status"]
    assert status == "pending"


def test_single_question_upload_is_unchanged(client, app, auth):
    auth.login()
    response = client.post(
        "/upload",
        data={
            "file": (io.BytesIO(b"student_id,response\nS001,x\n"), "a.csv"),
            "ideal_answer": "x",
            "max_score": "10",
        },
        content_type="multipart/form-data",
    )
    assert "/results/" in response.headers["Location"]
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) c FROM question").fetchone()["c"] == 0


# --- upgrade of an existing database ---------------------------------------

def test_existing_database_gains_question_support(tmp_path):
    path = str(tmp_path / "old.sqlite")
    conn = sqlite3.connect(path)
    conn.executescript(
        """CREATE TABLE instructor (id INTEGER PRIMARY KEY, email TEXT,
                                    password_hash TEXT, full_name TEXT);
           CREATE TABLE assignment (id INTEGER PRIMARY KEY);
           CREATE TABLE student_response (id INTEGER PRIMARY KEY, assignment_id INTEGER,
                                          student_id TEXT, response_text TEXT,
                                          row_number INTEGER);
           INSERT INTO instructor (email, password_hash) VALUES ('a@b.c', 'h');"""
    )
    conn.close()

    ensure_schema_current(path)
    ensure_schema_current(path)  # safe to run on every start

    conn = sqlite3.connect(path)
    columns = {r[1] for r in conn.execute("PRAGMA table_info(student_response)")}
    tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert conn.execute("SELECT COUNT(*) FROM instructor").fetchone()[0] == 1
    conn.close()
    assert "question_id" in columns
    assert "question" in tables
