"""Tests for the Excel export (UC-5, SS 2.23 right-to-left rendering, SS 2.24)."""
import io

from openpyxl import load_workbook

from qayyem.export import build_workbook, filename_for

ARABIC_FEEDBACK = (
    "إجابة جيدة "
    "وتحتاج إلى تفصيل."
)


class Row(dict):
    """A stand-in for a sqlite3.Row, which supports mapping access."""


def assignment(**overrides):
    base = {
        "id": 1,
        "title": "midterm.csv",
        "created_at": "2026-09-26 10:00:00",
        "max_score": 10,
        "ideal_answer": "The ideal answer.",
        "rubric": "Accuracy.",
    }
    base.update(overrides)
    return Row(base)


def results():
    return [
        Row({"student_id": "S001", "score": 8.5, "feedback_ar": ARABIC_FEEDBACK,
             "model_name": "gemini-2.0-flash", "error": None}),
        Row({"student_id": "S002", "score": 0, "feedback_ar": ARABIC_FEEDBACK,
             "model_name": "no-response", "error": None}),
        Row({"student_id": "S003", "score": None, "feedback_ar": "",
             "model_name": None, "error": "Grading failed: timeout"}),
    ]


def load(payload):
    return load_workbook(io.BytesIO(payload))


def test_workbook_has_results_and_metadata_sheets():
    book = load(build_workbook(assignment(), results()))
    assert book.sheetnames == ["Results", "Assignment"]


def test_headers_are_present():
    sheet = load(build_workbook(assignment(), results()))["Results"]
    assert [sheet.cell(row=1, column=c).value for c in range(1, 5)] == [
        "Student ID", "Score", "Feedback (Arabic)", "Note"
    ]


def test_every_student_appears_with_its_score():
    sheet = load(build_workbook(assignment(), results()))["Results"]
    assert sheet["A2"].value == "S001"
    assert sheet["B2"].value == 8.5
    assert sheet["A3"].value == "S002"
    assert sheet["B3"].value == 0


def test_arabic_feedback_is_written_right_to_left():
    """SS 2.23: the feedback column must render in right-to-left script."""
    sheet = load(build_workbook(assignment(), results()))["Results"]
    cell = sheet["C2"]
    assert cell.value == ARABIC_FEEDBACK
    assert cell.alignment.horizontal == "right"
    assert cell.alignment.readingOrder == 2
    assert cell.alignment.wrap_text is True


def test_ungraded_student_shows_dash_and_reason():
    sheet = load(build_workbook(assignment(), results()))["Results"]
    assert sheet["B4"].value == "-"
    assert "timeout" in sheet["D4"].value


def test_header_row_is_frozen():
    sheet = load(build_workbook(assignment(), results()))["Results"]
    assert sheet.freeze_panes == "A2"


def test_metadata_records_model_and_average():
    sheet = load(build_workbook(assignment(), results()))["Assignment"]
    labels = {sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).value
              for r in range(1, sheet.max_row + 1)}
    assert labels["Responses"] == 3
    assert labels["Graded successfully"] == 2
    assert labels["Average score"] == 4.25
    assert "gemini-2.0-flash" in labels["Model"]
    assert labels["Ideal answer"] == "The ideal answer."


def test_metadata_handles_no_graded_results():
    rows = [Row({"student_id": "S1", "score": None, "feedback_ar": "",
                 "model_name": None, "error": "failed"})]
    sheet = load(build_workbook(assignment(), rows))["Assignment"]
    labels = {sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).value
              for r in range(1, sheet.max_row + 1)}
    assert labels["Average score"] == "-"


def test_filename_is_safe_and_derived_from_title():
    assert filename_for(assignment(title="midterm.csv")) == "qayyem-midterm.xlsx"
    assert filename_for(assignment(title="Week 3 quiz.xlsx")) == "qayyem-Week-3-quiz.xlsx"
    unsafe = filename_for(assignment(title='a/b\\c:"*?<>|'))
    for char in '/\\:"*?<>|':
        assert char not in unsafe
    assert unsafe.endswith(".xlsx")
